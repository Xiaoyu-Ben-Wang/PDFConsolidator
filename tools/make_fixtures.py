"""Generate a mixed fixture folder for exercising the converter.

    python tools/make_fixtures.py [output_dir]
"""
from __future__ import annotations

import os
import sys

OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "fixtures", "generated")

os.makedirs(OUT, exist_ok=True)


def path(name: str) -> str:
    return os.path.join(OUT, name)


print("writing fixtures to %s" % OUT)

# ---------------------------------------------------------------- Word --------
from docx import Document
from docx.enum.text import WD_BREAK

doc = Document()
doc.add_heading("Quarterly Report", level=1)
for i in range(3):
    doc.add_paragraph("Paragraph %d. " % i + "Lorem ipsum dolor sit amet, " * 8)
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
doc.add_heading("Second page", level=2)
doc.add_paragraph("Content on the second page.")
doc.add_table(rows=3, cols=3)
for r in range(3):
    for c in range(3):
        doc.tables[0].cell(r, c).text = "R%dC%d" % (r + 1, c + 1)
doc.save(path("report.docx"))

doc2 = Document()
doc2.add_heading("Short memo", level=1)
doc2.add_paragraph("A single-page document with a little text.")
doc2.save(path("memo.docx"))

# --------------------------------------------------------------- Excel --------
from openpyxl import Workbook

wb = Workbook()
ws = wb.active
ws.title = "Budget"
ws.append(["Item", "Q1", "Q2", "Q3", "Q4", "Total"])
for i in range(1, 26):
    ws.append(["Item %02d" % i, i * 10, i * 11, i * 12, i * 13, "=SUM(B%d:E%d)" % (i + 1, i + 1)])
ws2 = wb.create_sheet("Notes")
ws2["A1"] = "Second sheet"
ws2["A2"] = "Tests multi-sheet export."
wb.save(path("budget.xlsx"))

# ---------------------------------------------------------- PowerPoint --------
from pptx import Presentation
from pptx.util import Inches

prs = Presentation()
for n in range(1, 4):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Slide %d" % n
    slide.placeholders[1].text = "Body text for slide %d" % n
    box = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(1))
    box.text_frame.text = "Extra textbox"
prs.save(path("deck.pptx"))

# --------------------------------------------------------------- images -------
from PIL import Image, ImageDraw

img = Image.new("RGB", (1600, 1200), (240, 244, 250))
d = ImageDraw.Draw(img)
d.rectangle([40, 40, 1560, 1160], outline=(40, 80, 160), width=6)
d.text((80, 80), "JPEG fixture 1600x1200", fill=(20, 20, 20))
img.save(path("photo.jpg"), quality=92)

# PNG with alpha, for the flatten-onto-white path in img2pdf.
alpha = Image.new("RGBA", (800, 600), (0, 0, 0, 0))
ad = ImageDraw.Draw(alpha)
ad.ellipse([100, 100, 700, 500], fill=(220, 60, 60, 255))
alpha.save(path("diagram.png"))

# Very tall image, for page fitting.
tall = Image.new("RGB", (600, 2400), (255, 255, 255))
td = ImageDraw.Draw(tall)
td.rectangle([20, 20, 580, 2380], outline=(0, 0, 0), width=4)
td.text((40, 40), "Tall scan 600x2400", fill=(0, 0, 0))
tall.save(path("scan_long.png"))

# ---------------------------------------------------------------- text --------
with open(path("notes.txt"), "w", encoding="utf-8") as fh:
    fh.write("Plain text fixture\n")
    fh.write("=" * 40 + "\n")
    for i in range(1, 21):
        fh.write("line %02d: some plain text content\n" % i)

with open(path("readme.md"), "w", encoding="utf-8") as fh:
    fh.write("# Markdown fixture\n\n")
    fh.write("This must render as *literal text*, not as Markdown.\n\n")
    fh.write("- bullet one\n- bullet two\n\n")
    fh.write("## A heading that should stay visible as '# heading'\n")

# ----------------------------------------------------------- existing PDF ----
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

c = canvas.Canvas(path("existing.pdf"), pagesize=A4)
for n in range(1, 3):
    c.setFont("Helvetica", 24)
    c.drawString(72, 760, "Existing PDF - page %d" % n)
    c.setFont("Helvetica", 11)
    c.drawString(72, 730, "A4 sized source, to exercise Letter normalisation.")
    c.showPage()
c.save()

# ------------------------------------------------------------------ report ----
print()
for name in sorted(os.listdir(OUT)):
    full = path(name)
    print("  %-18s %8d bytes" % (name, os.path.getsize(full)))
print("\n%d fixtures" % len(os.listdir(OUT)))
